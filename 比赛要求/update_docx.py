import os
import docx

# Find the docx file in the current directory
docx_files = [f for f in os.listdir('.') if f.endswith('.docx') and not f.startswith('~')]
if not docx_files:
    print('No docx file found!')
    exit(1)

path = docx_files[0]
print(f'Found document: {path}')

doc = docx.Document(path)
table = doc.tables[0]

# Add Row 3 for Meshy.ai
table.cell(3, 1).text = 'Meshy.ai (网页端/API) 2026.05 - 2026.05'
table.cell(3, 2).text = '3D模型生成、前端WebGL展示素材制作、3D资产贴图与优化'
table.cell(3, 3).text = '生成一个科幻风格的3D徽章模型，要求带有金属光泽和赛博朋克风的发光线条，导出为GLB格式。'
table.cell(3, 4).text = '快速生成了包含基础几何体、高质量PBR材质及贴图的3D模型文件。'
table.cell(3, 5).text = '下载GLB模型后，我们在前端使用Three.js调整了材质的金属度(metalness)和环境光反射，以更好地融入星空深色主题。'
table.cell(3, 6).text = '极大地降低了3D资产的制作门槛，为项目的数据可视化大屏和荣誉徽章界面提供了高质量的3D交互素材，节省了近一周的建模时间。'

# Check if there are other places needing updates
print('\n--- Checking Document Paragraphs ---')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f'Para {i}: {p.text}')

doc.save(path)
print('Document updated successfully with Meshy.ai!')
