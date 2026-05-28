import os
import docx

docx_files = [f for f in os.listdir('.') if f.endswith('.docx') and not f.startswith('~')]
if not docx_files:
    print('No docx file found!')
    exit(1)

path = docx_files[0]
doc = docx.Document(path)

# Replace placeholders in paragraphs
for p in doc.paragraphs:
    if '2026012345' in p.text:
        p.text = p.text.replace('2026012345', '202612010035')

# Also check tables if there are any other placeholders
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            if '2026012345' in cell.text:
                cell.text = cell.text.replace('2026012345', '202612010035')

doc.save(path)
print('Placeholders updated successfully!')
