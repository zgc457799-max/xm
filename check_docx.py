import docx

path = r'C:\Users\张广川\Desktop\xm\比赛要求\4-AI工具使用说明（选用模板）（2026年版）(4).docx'
doc = docx.Document(path)
table = doc.tables[0]

print('--- Table Contents ---')
for i, row in enumerate(table.rows):
    row_text = []
    for cell in row.cells:
        row_text.append(cell.text.replace('\n', ' '))
    print(f'Row {i}: {row_text}')

print('\n--- Document Paragraphs ---')
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f'Para {i}: {p.text}')
